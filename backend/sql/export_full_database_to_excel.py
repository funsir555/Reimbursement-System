#!/usr/bin/env python3
"""
Export full finex_db table data to a single .xlsx workbook.

Default workbook layout:
    - one worksheet per table
    - row 1: database column names
    - row 2: Chinese column comments
    - row 3+: table data

Large tables are split into multiple worksheets when the Excel row limit is reached.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from excel_mysql_common import (
    DEFAULT_APP_YML,
    ColumnMeta,
    ExcelToolError,
    fetch_columns,
    open_connection,
    quote_identifier,
    split_csv_arg,
)

DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / "templates"
EXCEL_MAX_ROWS = 1_048_576
INVALID_SHEET_TITLE_CHARS = re.compile(r'[\[\]\*:/\\?]')


@dataclass(frozen=True)
class TableExportSummary:
    table_name: str
    column_count: int
    row_count: int
    sheet_count: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export all MySQL base tables into one .xlsx workbook."
    )
    parser.add_argument(
        "--output",
        help=(
            "Output .xlsx path. Default: backend/sql/templates/"
            "finex_db_full_export_<timestamp>.xlsx"
        ),
    )
    parser.add_argument(
        "--app-yml",
        default=str(DEFAULT_APP_YML),
        help="Path to application.yml. Default: backend auth-service datasource config.",
    )
    parser.add_argument(
        "--tables",
        default="",
        help="Optional comma-separated table list. Default: export all base tables.",
    )
    parser.add_argument(
        "--exclude-tables",
        default="",
        help="Optional comma-separated table list to exclude.",
    )
    include_group = parser.add_mutually_exclusive_group()
    include_group.add_argument(
        "--include-comments",
        dest="include_comments",
        action="store_true",
        default=True,
        help="Include row 2 column comments. Default: enabled.",
    )
    include_group.add_argument(
        "--no-include-comments",
        dest="include_comments",
        action="store_false",
        help="Export only row 1 headers and row 2+ data.",
    )
    return parser.parse_args()


def ensure_dependencies() -> tuple[Any, Any]:
    try:
        import openpyxl  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit("Missing dependency 'openpyxl'. Run: pip install openpyxl pymysql") from exc
    try:
        import pymysql  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit("Missing dependency 'pymysql'. Run: pip install openpyxl pymysql") from exc
    return openpyxl, pymysql


def resolve_output_path(output_arg: str | None) -> Path:
    if output_arg:
        output_path = Path(output_arg).expanduser().resolve()
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = (DEFAULT_OUTPUT_DIR / f"finex_db_full_export_{timestamp}.xlsx").resolve()
    if output_path.suffix.lower() != ".xlsx":
        raise ExcelToolError("--output must end with .xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    return output_path


def normalize_table_list(raw: str) -> list[str]:
    names = split_csv_arg(raw)
    duplicates = [name for idx, name in enumerate(names) if name in names[:idx]]
    if duplicates:
        raise ExcelToolError(f"Duplicate table names found: {', '.join(duplicates)}")
    return names


def list_base_tables(cursor: Any) -> list[str]:
    cursor.execute(
        """
        SELECT TABLE_NAME
        FROM information_schema.TABLES
        WHERE TABLE_SCHEMA = DATABASE()
          AND TABLE_TYPE = 'BASE TABLE'
        ORDER BY TABLE_NAME
        """
    )
    return [row["TABLE_NAME"] for row in cursor.fetchall()]


def resolve_target_tables(cursor: Any, include_tables: list[str], exclude_tables: list[str]) -> list[str]:
    all_tables = list_base_tables(cursor)
    all_table_set = set(all_tables)

    missing_include = [name for name in include_tables if name not in all_table_set]
    if missing_include:
        raise ExcelToolError("Target table does not exist: " + ", ".join(missing_include))

    missing_exclude = [name for name in exclude_tables if name not in all_table_set]
    if missing_exclude:
        raise ExcelToolError("Excluded table does not exist: " + ", ".join(missing_exclude))

    exclude_set = set(exclude_tables)
    if include_tables:
        tables = [name for name in include_tables if name not in exclude_set]
    else:
        tables = [name for name in all_tables if name not in exclude_set]

    if not tables:
        raise ExcelToolError("No tables matched the export condition.")
    return tables


def clean_cell_value(openpyxl: Any, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bytes):
        return "0x" + value.hex()
    if isinstance(value, str):
        return openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def sanitize_sheet_title(raw: str) -> str:
    cleaned = INVALID_SHEET_TITLE_CHARS.sub("_", raw).strip().strip("'")
    return cleaned or "Sheet"


def allocate_sheet_title(base_name: str, part_index: int, used_titles: set[str]) -> str:
    suffix = "" if part_index == 1 else f"__{part_index}"
    cleaned = sanitize_sheet_title(base_name)
    title = (cleaned[: 31 - len(suffix)] + suffix)[:31]
    if title not in used_titles:
        used_titles.add(title)
        return title

    collision_index = 2
    while True:
        extra_suffix = f"_{collision_index}"
        core_limit = 31 - len(suffix) - len(extra_suffix)
        candidate = (cleaned[:core_limit] + suffix + extra_suffix)[:31]
        if candidate not in used_titles:
            used_titles.add(candidate)
            return candidate
        collision_index += 1


def create_table_sheet(
    workbook: Any,
    used_titles: set[str],
    table_name: str,
    part_index: int,
    headers: list[str],
    comments: list[str],
    include_comments: bool,
) -> Any:
    sheet_title = allocate_sheet_title(table_name, part_index, used_titles)
    sheet = workbook.create_sheet(title=sheet_title)
    sheet.append(headers)
    if include_comments:
        sheet.append(comments)
    return sheet


def export_single_table(
    connection: Any,
    pymysql: Any,
    openpyxl: Any,
    workbook: Any,
    used_titles: set[str],
    table_name: str,
    columns: list[ColumnMeta],
    include_comments: bool,
) -> TableExportSummary:
    headers = [column.name for column in columns]
    comments = [column.comment or "" for column in columns]
    meta_rows = 2 if include_comments else 1
    max_data_rows_per_sheet = EXCEL_MAX_ROWS - meta_rows

    total_rows = 0
    sheet_count = 1
    current_sheet = create_table_sheet(
        workbook=workbook,
        used_titles=used_titles,
        table_name=table_name,
        part_index=sheet_count,
        headers=headers,
        comments=comments,
        include_comments=include_comments,
    )
    current_sheet_data_rows = 0

    sql = f"SELECT * FROM {quote_identifier(table_name)}"
    with connection.cursor(pymysql.cursors.SSDictCursor) as cursor:
        cursor.execute(sql)
        while True:
            rows = cursor.fetchmany(500)
            if not rows:
                break
            for row in rows:
                if current_sheet_data_rows >= max_data_rows_per_sheet:
                    sheet_count += 1
                    current_sheet = create_table_sheet(
                        workbook=workbook,
                        used_titles=used_titles,
                        table_name=table_name,
                        part_index=sheet_count,
                        headers=headers,
                        comments=comments,
                        include_comments=include_comments,
                    )
                    current_sheet_data_rows = 0

                current_sheet.append(
                    [clean_cell_value(openpyxl, row.get(column.name)) for column in columns]
                )
                total_rows += 1
                current_sheet_data_rows += 1

    return TableExportSummary(
        table_name=table_name,
        column_count=len(columns),
        row_count=total_rows,
        sheet_count=sheet_count,
    )


def append_readme_sheet(
    workbook: Any,
    used_titles: set[str],
    output_path: Path,
    database_name: str,
    include_comments: bool,
    summaries: list[TableExportSummary],
) -> None:
    sheet = workbook.create_sheet(title=allocate_sheet_title("README", 1, used_titles))
    export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sheet.append(["导出说明", "值"])
    sheet.append(["导出时间", export_time])
    sheet.append(["数据库", database_name])
    sheet.append(["输出文件", str(output_path)])
    sheet.append(["是否包含注释行", "是" if include_comments else "否"])
    sheet.append(["导出表数量", len(summaries)])
    sheet.append(["备注", "每张表一个工作表；二进制列按十六进制字符串导出。"])
    sheet.append([])
    sheet.append(["表名", "列数", "数据行数", "工作表数"])
    for summary in summaries:
        sheet.append(
            [
                summary.table_name,
                summary.column_count,
                summary.row_count,
                summary.sheet_count,
            ]
        )


def parse_database_name(app_yml_path: Path) -> str:
    _, pymysql = ensure_dependencies()
    connection = None
    try:
        connection = open_connection(pymysql, app_yml_path)
        with connection.cursor() as cursor:
            cursor.execute("SELECT DATABASE() AS db_name")
            row = cursor.fetchone()
        database_name = row["db_name"] if row and row.get("db_name") else ""
        if not database_name:
            raise ExcelToolError("Failed to resolve current database name.")
        return str(database_name)
    finally:
        if connection is not None:
            connection.close()


def main() -> int:
    args = parse_args()
    openpyxl, pymysql = ensure_dependencies()
    app_yml_path = Path(args.app_yml).expanduser().resolve()
    output_path = resolve_output_path(args.output)
    include_tables = normalize_table_list(args.tables)
    exclude_tables = normalize_table_list(args.exclude_tables)

    connection = None
    try:
        connection = open_connection(pymysql, app_yml_path)
        with connection.cursor() as cursor:
            tables = resolve_target_tables(cursor, include_tables, exclude_tables)

        workbook = openpyxl.Workbook(write_only=True)
        used_titles: set[str] = set()
        summaries: list[TableExportSummary] = []

        print(f"[INFO] Output file: {output_path}")
        print(f"[INFO] Tables to export: {len(tables)}")

        for index, table_name in enumerate(tables, start=1):
            with connection.cursor() as cursor:
                column_map = fetch_columns(cursor, table_name)
            summary = export_single_table(
                connection=connection,
                pymysql=pymysql,
                openpyxl=openpyxl,
                workbook=workbook,
                used_titles=used_titles,
                table_name=table_name,
                columns=list(column_map.values()),
                include_comments=args.include_comments,
            )
            summaries.append(summary)
            print(
                f"[INFO] ({index}/{len(tables)}) {table_name}: "
                f"{summary.row_count} rows, {summary.column_count} columns, {summary.sheet_count} sheet(s)"
            )

        database_name = parse_database_name(app_yml_path)
        append_readme_sheet(
            workbook=workbook,
            used_titles=used_titles,
            output_path=output_path,
            database_name=database_name,
            include_comments=args.include_comments,
            summaries=summaries,
        )
        workbook.save(output_path)
        print(f"[INFO] Export finished: {output_path}")
        return 0
    except ExcelToolError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1
    except Exception as exc:  # pragma: no cover - defensive CLI fallback
        print(f"[ERROR] Unexpected failure: {exc}", file=sys.stderr)
        return 1
    finally:
        if connection is not None:
            connection.close()


if __name__ == "__main__":
    raise SystemExit(main())
